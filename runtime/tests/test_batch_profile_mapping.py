from pathlib import Path
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.db import Base
from app.models import ProfileRecord, WorkflowTemplateRecord
from app.services.batch_service import BatchService
from app.services.workflow_service import WorkflowService


class DummyProviderService:
    async def filter_managed_profiles(self, _session, _provider_type, profiles):
        return profiles


def workflow_payload() -> dict:
    return {
        "metadata": {"name": "Profile mapping workflow", "version": "1.0.0"},
        "profile_policy": {
            "provider_type": "ixbrowser",
            "selection_mode": "by_group",
            "group_ids": ["g1"],
            "profile_ids": [],
            "tag_ids": [],
        },
        "runtime_policy": {"mode": "visual"},
        "steps": [{"id": "goto-1", "type": "goto", "url": "https://example.com"}],
        "locators": {},
    }


def make_service(async_session) -> BatchService:
    workflow_service = WorkflowService(root_dir=Path("."))
    return BatchService(
        session_factory=async_session,
        registry=None,
        provider_service=DummyProviderService(),
        workflow_service=workflow_service,
        execution_service=None,
        monitor=None,
    )


async def seed_profiles_and_workflow(session) -> WorkflowTemplateRecord:
    workflow = WorkflowTemplateRecord(
        name="Profile mapping workflow",
        version="1.0.0",
        folder="未分组",
        target_provider_type="ixbrowser",
        workflow_yaml="",
        normalized_workflow_json=workflow_payload(),
        selector_catalog_json={},
    )
    session.add(workflow)
    session.add_all(
        [
            ProfileRecord(
                provider_type="ixbrowser",
                external_profile_id="101",
                display_name="Profile 101",
                group_summary={"id": "g1", "name": "Group 1"},
            ),
            ProfileRecord(
                provider_type="ixbrowser",
                external_profile_id="102",
                display_name="Profile 102",
                group_summary={"id": "g1", "name": "Group 1"},
            ),
        ]
    )
    await session.commit()
    await session.refresh(workflow)
    return workflow


@pytest.mark.asyncio
async def test_validate_profile_mapping_requires_exact_profile_id_match() -> None:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async_session = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    service = make_service(async_session)
    async with async_session() as session:
        workflow = await seed_profiles_and_workflow(session)

        valid = await service.validate_profile_mapping(
            session,
            workflow_id=workflow.id,
            provider_type="ixbrowser",
            rows=[{"profile_id": "101"}, {"profile_id": "102"}],
            strict=True,
        )
        assert valid.valid is True
        assert valid.matched_count == 2

        invalid = await service.validate_profile_mapping(
            session,
            workflow_id=workflow.id,
            provider_type="ixbrowser",
            rows=[{"profile_id": "101"}, {"profile_id": "101"}],
            strict=True,
        )
        assert invalid.valid is False
        assert invalid.duplicate_profile_ids == ["101"]
        assert invalid.missing_profile_ids == ["102"]


@pytest.mark.asyncio
async def test_build_input_template_contains_all_workflow_profiles() -> None:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async_session = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    service = make_service(async_session)
    async with async_session() as session:
        workflow = await seed_profiles_and_workflow(session)

        content = await service.build_input_template_xlsx(session, workflow_id=workflow.id)
        workbook = load_workbook(filename=BytesIO(content))
        assert "使用说明" in workbook.sheetnames
        rows = list(workbook.active.iter_rows(values_only=True))
        guide_text = "\n".join(
            str(cell or "")
            for row in workbook["使用说明"].iter_rows(values_only=True)
            for cell in row
        )

        assert rows[0][:7] == ("profile_id", "profile_name", "group_name", "profile_remark", "keyword", "search_keyword", "note")
        assert [row[0] for row in rows[1:]] == ["101", "102"]
        assert "Nike|Adidas|Puma" in guide_text
        assert "${row.search_keyword}" in guide_text
        assert "\\|" in guide_text
