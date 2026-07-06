from __future__ import annotations

import asyncio
import math
import platform
import re
import subprocess
import csv
import io
from dataclasses import dataclass, field
from datetime import datetime
from itertools import cycle
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.config import settings
from app.models import (
    BatchRecord,
    BatchRowRecord,
    BatchStatus,
    ProfileRecord,
    StepRunRecord,
    TaskRunRecord,
    TaskRunStatus,
)
from app.providers.registry import ProviderRegistry
from app.schemas.batch import BatchDetail, BatchImportResult, BatchRowPayload, BatchSummary, StartBatchRequest, TaskRunDetail, StepRunDetail
from app.schemas.provider import ProviderSessionRef
from app.schemas.workflow import WorkflowDefinition
from app.services.execution_service import ExecutionService
from app.services.monitor_service import MonitorService
from app.services.provider_service import ProviderService
from app.services.workflow_service import WorkflowService


@dataclass
class BatchControl:
    paused: asyncio.Event = field(default_factory=asyncio.Event)
    cancelled: bool = False

    def __post_init__(self) -> None:
        self.paused.set()


class BatchService:
    def __init__(
        self,
        *,
        session_factory: async_sessionmaker[AsyncSession],
        registry: ProviderRegistry,
        provider_service: ProviderService,
        workflow_service: WorkflowService,
        execution_service: ExecutionService,
        monitor: MonitorService,
    ) -> None:
        self.session_factory = session_factory
        self.registry = registry
        self.provider_service = provider_service
        self.workflow_service = workflow_service
        self.execution_service = execution_service
        self.monitor = monitor
        self._active_jobs: dict[str, asyncio.Task[None]] = {}
        self._controls: dict[str, BatchControl] = {}

    async def list_batches(self, session: AsyncSession) -> list[BatchSummary]:
        result = await session.execute(select(BatchRecord).order_by(BatchRecord.created_at.desc()))
        batches = []
        for row in result.scalars():
            batches.append(self._to_summary(row))
        return batches

    async def import_batch(
        self,
        session: AsyncSession,
        *,
        file_name: str,
        content: bytes,
        provider_type: str,
    ) -> BatchImportResult:
        rows = self._parse_input_file(file_name, content)
        batch = BatchRecord(
            name=f"Imported {file_name}",
            provider_type=provider_type,
            workflow_id="",
            status=BatchStatus.DRAFT,
            input_source="file",
            total_rows=len(rows),
        )
        session.add(batch)
        await session.flush()
        for index, payload in enumerate(rows, start=1):
            session.add(
                BatchRowRecord(
                    batch_id=batch.id,
                    row_index=index,
                    row_payload_json=payload,
                    dedupe_key=str(payload.get("id") or payload.get("name") or index),
                )
            )
        await session.commit()
        await session.refresh(batch)
        return BatchImportResult(
            batch=await self.get_batch_detail(session, batch.id),
            detected_columns=list(rows[0].keys()) if rows else [],
            preview_rows=rows[:5],
        )

    async def get_batch_detail(self, session: AsyncSession, batch_id: str) -> BatchDetail:
        batch = await session.get(BatchRecord, batch_id)
        if not batch:
            raise ValueError(f"batch not found: {batch_id}")
        row_result = await session.execute(
            select(BatchRowRecord).where(BatchRowRecord.batch_id == batch_id).order_by(BatchRowRecord.row_index.asc())
        )
        rows = [
            BatchRowPayload(row_index=row.row_index, dedupe_key=row.dedupe_key, payload=row.row_payload_json)
            for row in row_result.scalars()
        ]
        return BatchDetail(
            **self._to_summary(batch).model_dump(),
            profile_policy_snapshot=batch.profile_policy_snapshot or {},
            result_summary_json=batch.result_summary_json or {},
            rows=rows,
        )

    async def start_batch(self, session: AsyncSession, batch_id: str, request: StartBatchRequest) -> BatchDetail:
        batch = await session.get(BatchRecord, batch_id)
        if not batch:
            raise ValueError(f"batch not found: {batch_id}")
        batch.workflow_id = request.workflow_id
        batch.provider_type = request.provider_type
        batch.runtime_mode = request.runtime_mode
        batch.profile_policy_snapshot = request.profile_policy_snapshot
        batch.status = BatchStatus.READY
        await session.commit()

        control = BatchControl()
        self._controls[batch.id] = control
        task = asyncio.create_task(self._run_batch(batch.id, request.requested_slots, control))
        self._active_jobs[batch.id] = task
        return await self.get_batch_detail(session, batch.id)

    async def pause_batch(self, session: AsyncSession, batch_id: str) -> BatchSummary:
        control = self._controls.setdefault(batch_id, BatchControl())
        control.paused.clear()
        batch = await session.get(BatchRecord, batch_id)
        if batch:
            batch.status = BatchStatus.PAUSED
            await session.commit()
            return self._to_summary(batch)
        raise ValueError(f"batch not found: {batch_id}")

    async def resume_batch(self, session: AsyncSession, batch_id: str) -> BatchSummary:
        control = self._controls.setdefault(batch_id, BatchControl())
        control.paused.set()
        batch = await session.get(BatchRecord, batch_id)
        if batch:
            batch.status = BatchStatus.RUNNING
            await session.commit()
            return self._to_summary(batch)
        raise ValueError(f"batch not found: {batch_id}")

    async def cancel_batch(self, session: AsyncSession, batch_id: str) -> BatchSummary:
        control = self._controls.setdefault(batch_id, BatchControl())
        control.cancelled = True
        control.paused.set()
        batch = await session.get(BatchRecord, batch_id)
        if batch:
            batch.status = BatchStatus.CANCELLED
            await session.commit()
            return self._to_summary(batch)
        raise ValueError(f"batch not found: {batch_id}")

    async def get_task_detail(self, session: AsyncSession, task_run_id: str) -> TaskRunDetail:
        task = await session.get(TaskRunRecord, task_run_id)
        if not task:
            raise ValueError(f"task not found: {task_run_id}")
        step_result = await session.execute(
            select(StepRunRecord).where(StepRunRecord.task_run_id == task_run_id).order_by(StepRunRecord.created_at.asc())
        )
        steps = [
            StepRunDetail(
                id=row.id,
                step_id=row.step_id,
                label=row.label,
                action_type=row.action_type,
                status=row.status.value,
                error_code=row.error_code,
                error_message=row.error_message,
                output_payload_json=row.output_payload_json or {},
                locator_summary_json=row.locator_summary_json or {},
                artifact_path=row.artifact_path,
            )
            for row in step_result.scalars()
        ]
        return TaskRunDetail(
            id=task.id,
            batch_id=task.batch_id,
            provider_type=task.provider_type,
            provider_profile_id=task.provider_profile_id,
            status=task.status.value,
            slot_index=task.slot_index,
            error_code=task.error_code,
            error_message=task.error_message,
            outputs_json=task.outputs_json or {},
            step_runs=steps,
        )

    async def _run_batch(self, batch_id: str, requested_slots: int, control: BatchControl) -> None:
        async with self.session_factory() as session:
            batch = await session.get(BatchRecord, batch_id)
            if not batch:
                return
            workflow_record = await self.workflow_service.get_workflow(session, batch.workflow_id)
            workflow = WorkflowDefinition.model_validate(workflow_record.normalized_workflow_json)
            profiles, empty_reason = await self._resolve_profiles(session, batch.provider_type, batch.profile_policy_snapshot or {})
            rows_result = await session.execute(
                select(BatchRowRecord).where(BatchRowRecord.batch_id == batch_id).order_by(BatchRowRecord.row_index.asc())
            )
            rows = list(rows_result.scalars())
            if not rows:
                batch.status = BatchStatus.COMPLETED
                batch.result_summary_json = {"total": 0, "success": 0, "failed": 0}
                await session.commit()
                return
            if not profiles:
                batch.status = BatchStatus.FAILED
                batch.result_summary_json = {
                    "error": "Provider 管理范围未命中任何 Profile"
                    if empty_reason == "provider_scope_empty"
                    else "No cached profiles matched the batch profile policy"
                }
                await session.commit()
                return
            screen_bounds = await self._read_screen_bounds()
            screen_capacity = self._estimate_visual_slot_capacity(screen_bounds)
            slot_limit = max(1, min(requested_slots, settings.max_slot_limit, screen_capacity, len(profiles), len(rows)))
            batch.status = BatchStatus.RUNNING
            await session.commit()
            assignments = cycle(profiles)
            tasks: list[TaskRunRecord] = []
            for row in rows:
                profile = next(assignments)
                task_run = TaskRunRecord(
                    batch_id=batch.id,
                    batch_row_id=row.id,
                    provider_type=batch.provider_type,
                    provider_profile_id=profile.external_profile_id,
                    slot_index=None,
                )
                session.add(task_run)
                tasks.append(task_run)
            await session.commit()

        provider = self.registry.get(batch.provider_type)
        task_queue: asyncio.Queue[str] = asyncio.Queue()
        for task in tasks:
            task_queue.put_nowait(task.id)

        active_sessions: dict[int, ProviderSessionRef] = {}
        layout_lock = asyncio.Lock()

        async def arrange_active_windows() -> None:
            async with layout_lock:
                sessions = [active_sessions[index] for index in sorted(active_sessions)]
                if not sessions:
                    return
                layout = self._build_provider_tile_layout(
                    active_count=len(sessions),
                    slot_limit=slot_limit,
                    runtime_policy=workflow.runtime_policy.model_dump(),
                    screen_bounds=screen_bounds,
                )
                native_ok = True
                native_error = None
                try:
                    await provider.arrange_windows(layout)
                except Exception as exc:  # noqa: BLE001
                    native_ok = False
                    native_error = str(exc)
                    await self._arrange_windows_macos(
                        sessions=sessions,
                        slot_limit=slot_limit,
                        runtime_policy=workflow.runtime_policy.model_dump(),
                        screen_bounds=screen_bounds,
                    )
                await self.monitor.publish(
                    "batch.layout",
                    {
                        "batch_id": batch_id,
                        "active_count": len(sessions),
                        "slot_limit": slot_limit,
                        "native_provider_layout": native_ok,
                        "native_error": native_error,
                    },
                )

        async def worker(slot_index: int) -> None:
            while True:
                await control.paused.wait()
                if control.cancelled:
                    return
                try:
                    task_run_id = task_queue.get_nowait()
                except asyncio.QueueEmpty:
                    return

                async def on_session_opened(_: str, session_ref: ProviderSessionRef) -> None:
                    active_sessions[slot_index] = session_ref
                    await arrange_active_windows()

                async def on_session_closed(_: str, __: str) -> None:
                    active_sessions.pop(slot_index, None)
                    await arrange_active_windows()

                try:
                    await self._assign_task_slot(task_run_id, slot_index)
                    await self.execution_service.execute_task(
                        task_run_id=task_run_id,
                        workflow=workflow,
                        provider=provider,
                        on_session_opened=on_session_opened,
                        on_session_closed=on_session_closed,
                    )
                    await self.monitor.publish("batch.progress", {"batch_id": batch_id, "task_run_id": task_run_id})
                except Exception as exc:  # noqa: BLE001
                    await self._mark_task_worker_failed(task_run_id, exc)
                finally:
                    task_queue.task_done()

        try:
            await asyncio.gather(*(worker(slot_index) for slot_index in range(slot_limit)))
            await self._mark_unfinished_tasks(batch_id, cancelled=control.cancelled)
            await self._finalize_batch(batch_id, cancelled=control.cancelled)
        finally:
            self._active_jobs.pop(batch_id, None)
            self._controls.pop(batch_id, None)

    async def _assign_task_slot(self, task_run_id: str, slot_index: int) -> None:
        async with self.session_factory() as session:
            task = await session.get(TaskRunRecord, task_run_id)
            if task and task.status == TaskRunStatus.QUEUED:
                task.slot_index = slot_index
                await session.commit()

    async def _mark_task_worker_failed(self, task_run_id: str, exc: Exception) -> None:
        async with self.session_factory() as session:
            task = await session.get(TaskRunRecord, task_run_id)
            if task and task.status not in {TaskRunStatus.SUCCEEDED, TaskRunStatus.FAILED, TaskRunStatus.CANCELLED}:
                task.status = TaskRunStatus.FAILED
                task.error_code = "worker_exception"
                task.error_message = str(exc)
                task.finished_at = datetime.utcnow()
                await session.commit()

    async def _mark_unfinished_tasks(self, batch_id: str, *, cancelled: bool) -> None:
        terminal_statuses = {
            TaskRunStatus.SUCCEEDED,
            TaskRunStatus.FAILED,
            TaskRunStatus.CANCELLED,
            TaskRunStatus.LOST,
        }
        async with self.session_factory() as session:
            result = await session.execute(select(TaskRunRecord).where(TaskRunRecord.batch_id == batch_id))
            for task in result.scalars():
                if task.status in terminal_statuses:
                    continue
                task.status = TaskRunStatus.CANCELLED if cancelled else TaskRunStatus.LOST
                task.error_code = "batch_cancelled" if cancelled else "runtime_interrupted"
                task.error_message = "批次已取消" if cancelled else "任务未正常结束，已由批次收口标记为 LOST"
                task.finished_at = datetime.utcnow()
            await session.commit()

    async def _finalize_batch(self, batch_id: str, *, cancelled: bool) -> None:
        async with self.session_factory() as session:
            batch = await session.get(BatchRecord, batch_id)
            if not batch:
                return
            task_result = await session.execute(select(TaskRunRecord).where(TaskRunRecord.batch_id == batch_id))
            task_runs = list(task_result.scalars())
            success_count = sum(1 for task in task_runs if task.status == TaskRunStatus.SUCCEEDED)
            failed_count = sum(1 for task in task_runs if task.status == TaskRunStatus.FAILED)
            lost_count = sum(1 for task in task_runs if task.status == TaskRunStatus.LOST)
            cancelled_count = sum(1 for task in task_runs if task.status == TaskRunStatus.CANCELLED)
            failure_count = failed_count + lost_count
            batch.success_count = success_count
            batch.failure_count = failure_count
            if cancelled:
                batch.status = BatchStatus.CANCELLED
            else:
                batch.status = BatchStatus.COMPLETED if failure_count == 0 and cancelled_count == 0 else BatchStatus.FAILED
            batch.result_summary_json = {
                "total": len(task_runs),
                "success": success_count,
                "failed": failed_count,
                "lost": lost_count,
                "cancelled": cancelled_count,
                "slot_pool": True,
            }
            await session.commit()

    async def recover_interrupted_batches(self) -> None:
        async with self.session_factory() as session:
            result = await session.execute(
                select(BatchRecord).where(
                    BatchRecord.status.in_([BatchStatus.READY, BatchStatus.RUNNING, BatchStatus.PAUSED])
                )
            )
            batch_ids = [batch.id for batch in result.scalars()]
        for batch_id in batch_ids:
            await self._mark_unfinished_tasks(batch_id, cancelled=False)
            await self._finalize_batch(batch_id, cancelled=False)

    @staticmethod
    def _build_provider_tile_layout(
        *,
        active_count: int,
        slot_limit: int,
        runtime_policy: dict[str, Any],
        screen_bounds: tuple[int, int, int, int] | None = None,
    ) -> dict[str, int]:
        per_line = max(1, math.ceil(math.sqrt(max(active_count, slot_limit))))
        rows = max(1, math.ceil(max(active_count, slot_limit) / per_line))
        preferred_width = max(
            settings.window_layout_min_width,
            int(runtime_policy.get("min_window_width") or settings.window_layout_default_width),
        )
        preferred_height = max(
            settings.window_layout_min_height,
            int(runtime_policy.get("min_window_height") or settings.window_layout_default_height),
        )
        width = preferred_width
        height = preferred_height
        if screen_bounds:
            left, top, right, bottom = screen_bounds
            usable_width = max(1, right - left - (settings.window_layout_margin_px * (per_line + 1)))
            usable_height = max(
                1,
                bottom
                - top
                - settings.window_layout_bottom_reserved_px
                - (settings.window_layout_margin_px * (rows + 1)),
            )
            width = max(settings.window_layout_min_width, min(preferred_width, int(usable_width / per_line)))
            height = max(settings.window_layout_min_height, min(preferred_height, int(usable_height / rows)))
        return {
            "screen": settings.window_layout_screen_index,
            "layout": 1,
            "adaptive": 1,
            "starting_position_x": settings.window_layout_margin_px,
            "starting_position_y": settings.window_layout_margin_px,
            "profile_size_width": width,
            "profile_size_hight": height,
            "profile_spacing_horizontal": settings.window_layout_margin_px,
            "profile_spacing_vertical": settings.window_layout_margin_px,
            "profile_deviaton_x": settings.window_layout_provider_deviation_px,
            "profile_deviaton_y": settings.window_layout_provider_deviation_px,
            "per_line_number_of_profiles": per_line,
        }

    async def _arrange_windows_macos(
        self,
        *,
        sessions: list[ProviderSessionRef],
        slot_limit: int,
        runtime_policy: dict[str, Any],
        screen_bounds: tuple[int, int, int, int] | None = None,
    ) -> None:
        if platform.system() != "Darwin":
            return
        pids = [session.browser_pid for session in sessions if session.browser_pid]
        if not pids:
            return
        script = self._build_macos_layout_script(
            pids=pids,
            slot_limit=slot_limit,
            min_width=settings.window_layout_min_width,
            min_height=settings.window_layout_min_height,
            screen_bounds=screen_bounds,
        )
        try:
            await asyncio.to_thread(
                subprocess.run,
                ["osascript", "-e", script],
                capture_output=True,
                text=True,
                timeout=settings.window_layout_timeout_sec,
                check=False,
            )
        except Exception:  # noqa: BLE001
            return

    @staticmethod
    def _build_macos_layout_script(
        *,
        pids: list[int],
        slot_limit: int,
        min_width: int,
        min_height: int,
        screen_bounds: tuple[int, int, int, int] | None = None,
    ) -> str:
        pid_list = ", ".join(str(pid) for pid in pids)
        columns = max(1, math.ceil(math.sqrt(max(len(pids), slot_limit))))
        rows = max(1, math.ceil(max(len(pids), slot_limit) / columns))
        screen_init = ""
        if screen_bounds:
            left, top, right, bottom = screen_bounds
            screen_init = f"""
set screenLeft to {left}
set screenTop to {top}
set screenRight to {right}
set screenBottom to {bottom}
"""
        else:
            screen_init = """
tell application "Finder" to set screenBounds to bounds of window of desktop
set screenLeft to item 1 of screenBounds
set screenTop to item 2 of screenBounds
set screenRight to item 3 of screenBounds
set screenBottom to item 4 of screenBounds
"""
        return f"""
set targetPids to {{{pid_list}}}
set columnsCount to {columns}
set rowsCount to {rows}
set minWidth to {min_width}
set minHeight to {min_height}
set marginSize to {settings.window_layout_margin_px}
{screen_init}
set usableWidth to screenRight - screenLeft - (marginSize * (columnsCount + 1))
set usableHeight to screenBottom - screenTop - {settings.window_layout_bottom_reserved_px} - (marginSize * (rowsCount + 1))
set cellWidth to usableWidth / columnsCount
set cellHeight to usableHeight / rowsCount
if cellWidth < minWidth and ((minWidth * columnsCount) + (marginSize * (columnsCount + 1))) <= (screenRight - screenLeft) then set cellWidth to minWidth
if cellHeight < minHeight and ((minHeight * rowsCount) + (marginSize * (rowsCount + 1)) + {settings.window_layout_bottom_reserved_px}) <= (screenBottom - screenTop) then set cellHeight to minHeight
tell application "System Events"
  set targetWindows to {{}}
  repeat with p in processes
    try
      if targetPids contains (unix id of p) then
        if (count of windows of p) > 0 then set end of targetWindows to item 1 of windows of p
      end if
    end try
  end repeat
  repeat with windowIndex from 1 to count of targetWindows
    set zeroIndex to windowIndex - 1
    set columnIndex to zeroIndex mod columnsCount
    set rowIndex to zeroIndex div columnsCount
    set xPosition to screenLeft + marginSize + (columnIndex * (cellWidth + marginSize))
    set yPosition to screenTop + 30 + marginSize + (rowIndex * (cellHeight + marginSize))
    try
      set position of item windowIndex of targetWindows to {{xPosition, yPosition}}
      set size of item windowIndex of targetWindows to {{cellWidth, cellHeight}}
    end try
  end repeat
end tell
"""

    async def _read_screen_bounds(self) -> tuple[int, int, int, int] | None:
        if platform.system() != "Darwin":
            return None
        try:
            result = await asyncio.to_thread(
                subprocess.run,
                ["osascript", "-e", 'tell application "Finder" to get bounds of window of desktop'],
                capture_output=True,
                text=True,
                timeout=settings.window_layout_timeout_sec,
                check=False,
            )
        except Exception:  # noqa: BLE001
            return None
        if result.returncode != 0:
            return None
        numbers = [int(item) for item in re.findall(r"-?\d+", result.stdout)]
        if len(numbers) < 4:
            return None
        left, top, right, bottom = numbers[:4]
        if right <= left or bottom <= top:
            return None
        return left, top, right, bottom

    @staticmethod
    def _estimate_visual_slot_capacity(screen_bounds: tuple[int, int, int, int] | None) -> int:
        if not screen_bounds:
            return settings.max_slot_limit
        left, top, right, bottom = screen_bounds
        usable_width = max(1, right - left - settings.window_layout_margin_px)
        usable_height = max(
            1,
            bottom - top - settings.window_layout_bottom_reserved_px - settings.window_layout_margin_px,
        )
        columns = max(
            1,
            usable_width // (settings.window_layout_min_width + settings.window_layout_margin_px),
        )
        rows = max(
            1,
            usable_height // (settings.window_layout_min_height + settings.window_layout_margin_px),
        )
        return max(1, min(settings.max_slot_limit, int(columns * rows)))

    async def _resolve_profiles(
        self,
        session: AsyncSession,
        provider_type: str,
        policy: dict[str, Any],
    ) -> tuple[list[ProfileRecord], str | None]:
        result = await session.execute(
            select(ProfileRecord).where(ProfileRecord.provider_type == provider_type, ProfileRecord.enabled.is_(True))
        )
        profiles = await self.provider_service.filter_managed_profiles(
            session,
            provider_type,
            list(result.scalars()),
        )
        if not profiles:
            return [], "provider_scope_empty"
        selection_mode = policy.get("selection_mode", "explicit_profiles")
        profiles = self.apply_profile_policy(profiles, policy)
        if not profiles:
            return [], "profile_policy_empty"
        return profiles, None

    @staticmethod
    def apply_profile_policy(
        profiles: list[ProfileRecord],
        policy: dict[str, Any],
    ) -> list[ProfileRecord]:
        selection_mode = policy.get("selection_mode", "explicit_profiles")
        if selection_mode == "explicit_profiles" and policy.get("profile_ids"):
            ids = {str(item) for item in policy["profile_ids"]}
            profiles = [item for item in profiles if item.external_profile_id in ids]
        elif selection_mode == "by_group" and policy.get("group_ids"):
            ids = {str(item) for item in policy["group_ids"]}
            profiles = [item for item in profiles if str(item.group_summary.get("id")) in ids]
        elif selection_mode == "by_tag" and policy.get("tag_ids"):
            ids = {str(item) for item in policy["tag_ids"]}
            profiles = [
                item
                for item in profiles
                if any(str(tag.get("id")) in ids for tag in (item.tag_summary or []))
            ]
        return profiles

    @staticmethod
    def _parse_input_file(file_name: str, content: bytes) -> list[dict[str, Any]]:
        suffix = Path(file_name).suffix.lower()
        if suffix in {".csv", ".txt"}:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            return [dict(row) for row in reader]
        if suffix in {".xlsx", ".xlsm"}:
            workbook = load_workbook(io.BytesIO(content), read_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return []
            header = [str(item) for item in rows[0]]
            return [dict(zip(header, row, strict=False)) for row in rows[1:]]
        raise ValueError(f"unsupported input file: {suffix}")

    @staticmethod
    def _to_summary(batch: BatchRecord) -> BatchSummary:
        return BatchSummary(
            id=batch.id,
            name=batch.name,
            provider_type=batch.provider_type,
            workflow_id=batch.workflow_id,
            status=batch.status.value,
            runtime_mode=batch.runtime_mode,
            total_rows=batch.total_rows,
            success_count=batch.success_count,
            failure_count=batch.failure_count,
            average_duration_ms=batch.average_duration_ms,
            created_at=batch.created_at,
            updated_at=batch.updated_at,
        )
